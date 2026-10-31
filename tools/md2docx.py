#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown → DOCX（用 python-docx 直接构建，不经过 HTML 中转）

为什么不用 textutil：
    实测 textutil 由 HTML 转 DOCX 会把 <table> 压成制表符文本，
    表格边框、合并、列宽全部丢失。文档里有 20 多张表，不能接受。

为什么不用 calibre：
    calibre 的 DOCX 生成器产出的表格缺少 <w:tblGrid>，
    python-docx 无法解析，Word 打开也可能提示修复。

本脚本直接操作 docx 文档对象模型：真表格、真标题样式、A4 页面、
中文字体（正文宋体 / 标题黑体）、代码块底纹、列表缩进。
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = Path('/Users/bu/Documents/乡育心桥/docs/技术文档.md')
OUT = Path.home() / 'Desktop' / '乡育心桥_技术文档.docx'

SERIF = '宋体'
SANS = '黑体'


# ---------------- 基础工具 ----------------
def set_cn_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade(cell_or_par, hex_fill):
    """给单元格或段落加底色"""
    el = cell_or_par._element
    pr = el.find(qn('w:tcPr')) if el.tag.endswith('tc') else el.find(qn('w:pPr'))
    if pr is None:
        pr = OxmlElement('w:tcPr' if el.tag.endswith('tc') else 'w:pPr')
        el.insert(0, pr)
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), hex_fill)
    pr.append(sh)


def add_borders(table):
    """表格加全边框（python-docx 默认 Table Grid 样式已含，这里显式指定更保险）"""
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '6')
        e.set(qn('w:color'), 'B9C3CE')
        borders.append(e)
    tblPr.append(borders)


# ---------------- 行内格式 ----------------
INLINE_RE = re.compile(r'(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)')


def add_inline(par, text, base_font=SERIF, size=10.5, color=None):
    """把 markdown 行内语法写进段落"""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            set_cn_font(r, base_font, size, color=color)
        tok = m.group(0)
        if tok.startswith('`'):
            r = par.add_run(tok[1:-1])
            set_cn_font(r, 'Menlo', size - 1, color=(0xA3, 0x33, 0x33))
        elif tok.startswith('**'):
            r = par.add_run(tok[2:-2])
            set_cn_font(r, base_font, size, bold=True, color=color)
        else:
            r = par.add_run(tok[1:-1])
            set_cn_font(r, base_font, size, bold=False, color=color)
            r.font.italic = True
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:])
        set_cn_font(r, base_font, size, color=color)


# ---------------- 主转换 ----------------
def parse(md):
    """切成块：('h',level,text) ('p',text) ('ul',[items]) ('ol',[items])
                 ('table',header,rows) ('code',lines) ('hr',)"""
    lines = md.split('\n')
    blocks = []
    i, n = 0, len(lines)

    def is_sep(s):
        s = s.strip()
        return bool(re.fullmatch(r'\|?[\s:\-|]+\|?', s)) and '-' in s

    while i < n:
        raw = lines[i]
        s = raw.strip()

        if not s:
            i += 1
            continue

        # 代码块
        if s.startswith('```'):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(('code', buf))
            continue

        # 分隔线
        if re.fullmatch(r'-{3,}', s):
            blocks.append(('hr',))
            i += 1
            continue

        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            blocks.append(('h', len(m.group(1)), m.group(2)))
            i += 1
            continue

        # 引用
        if s.startswith('>'):
            buf = []
            while i < n and lines[i].strip().startswith('>'):
                buf.append(lines[i].strip().lstrip('>').strip())
                i += 1
            blocks.append(('quote', ' '.join(buf)))
            continue

        # 表格
        if '|' in s and i + 1 < n and is_sep(lines[i + 1]):
            header = [c.strip() for c in s.strip('|').split('|')]
            i += 2
            rows = []
            while i < n and '|' in lines[i] and lines[i].strip():
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            blocks.append(('table', header, rows))
            continue

        # 有序列表
        if re.match(r'^\d+\.\s+', s):
            items = []
            while i < n and re.match(r'^\d+\.\s+', lines[i].strip()):
                items.append(re.sub(r'^\d+\.\s+', '', lines[i].strip()))
                i += 1
            blocks.append(('ol', items))
            continue

        # 无序列表
        if re.match(r'^[-*]\s+', s):
            items = []
            while i < n and re.match(r'^[-*]\s+', lines[i].strip()):
                items.append(re.sub(r'^[-*]\s+', '', lines[i].strip()))
                i += 1
            blocks.append(('ul', items))
            continue

        # 段落
        buf = [s]
        i += 1
        while i < n and lines[i].strip() \
                and not re.match(r'^(#{1,6}\s|```|>|[-*]\s|\d+\.\s)', lines[i].strip()) \
                and '|' not in lines[i]:
            buf.append(lines[i].strip())
            i += 1
        blocks.append(('p', ' '.join(buf)))

    return blocks


def build(blocks):
    doc = Document()

    # 页面：A4 + 常规页边距
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    # 正文默认样式
    st = doc.styles['Normal']
    st.font.name = SERIF
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), SERIF)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    st.paragraph_format.line_spacing = 1.55
    st.paragraph_format.space_after = Pt(4)

    usable_cm = 21.0 - 2.0 - 2.0

    for b in blocks:
        kind = b[0]

        if kind == 'h':
            _, level, text = b
            p = doc.add_paragraph()
            pf = p.paragraph_format
            if level == 1:
                pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf.space_before = Pt(0)
                pf.space_after = Pt(14)
                add_inline(p, text, SANS, 19)
                for r in p.runs:
                    r.font.bold = True
                # 标题下加一条横线
                pPr = p._element.get_or_add_pPr()
                pbdr = OxmlElement('w:pBdr')
                bt = OxmlElement('w:bottom')
                bt.set(qn('w:val'), 'single')
                bt.set(qn('w:sz'), '12')
                bt.set(qn('w:color'), '1B5E8C')
                bt.set(qn('w:space'), '6')
                pbdr.append(bt)
                pPr.append(pbdr)
            elif level == 2:
                pf.space_before = Pt(18)
                pf.space_after = Pt(8)
                add_inline(p, text, SANS, 14)
                for r in p.runs:
                    r.font.bold = True
            elif level == 3:
                pf.space_before = Pt(13)
                pf.space_after = Pt(5)
                add_inline(p, text, SANS, 12)
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x14, 0x3B, 0x52)
            else:
                pf.space_before = Pt(10)
                pf.space_after = Pt(4)
                add_inline(p, text, SANS, 11)
                for r in p.runs:
                    r.font.bold = True

        elif kind == 'p':
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            add_inline(p, b[1])

        elif kind == 'quote':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, b[1], SERIF, 10, color=(0x3C, 0x4C, 0x5C))
            shade(p, 'F4F8FC')

        elif kind == 'ul':
            for it in b[1]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.4)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run('· ')
                set_cn_font(r, SERIF, 10.5)
                add_inline(p, it)

        elif kind == 'ol':
            for idx, it in enumerate(b[1], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.85)
                p.paragraph_format.first_line_indent = Cm(-0.55)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(f'{idx}. ')
                set_cn_font(r, SERIF, 10.5)
                add_inline(p, it)

        elif kind == 'code':
            for ln in b[1]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.3)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.25
                r = p.add_run(ln if ln.strip() else ' ')
                set_cn_font(r, 'Menlo', 9, color=(0x1A, 0x1A, 0x1A))
                shade(p, 'F5F7F9')

        elif kind == 'table':
            _, header, rows = b
            cols = len(header)
            t = doc.add_table(rows=1, cols=cols)
            t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = True
            add_borders(t)

            # 表头
            hdr = t.rows[0].cells
            for j, h in enumerate(header):
                cell = hdr[j]
                cell.text = ''
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                add_inline(p, h, SANS, 9.5)
                for r in p.runs:
                    r.font.bold = True
                shade(cell, 'EEF2F6')

            # 数据行
            for row in rows:
                cells = t.add_row().cells
                for j in range(cols):
                    val = row[j] if j < len(row) else ''
                    cell = cells[j]
                    cell.text = ''
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.3
                    add_inline(p, val, SERIF, 9.5)

            # 表格后留白
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

        elif kind == 'hr':
            pass

    doc.save(str(OUT))
    return OUT


def main():
    md = SRC.read_text(encoding='utf-8')
    blocks = parse(md)
    out = build(blocks)

    from docx import Document as D2
    d = D2(str(out))
    size_kb = out.stat().st_size // 1024
    print(f'生成 {out}')
    print(f'  段落 {len(d.paragraphs)} · 表格 {len(d.tables)} · {size_kb} KB')
    if d.tables:
        t = d.tables[0]
        print(f'  首表 {len(t.rows)} 行 x {len(t.columns)} 列，首行：'
              + ' | '.join(c.text[:10] for c in t.rows[0].cells))
    return 0


if __name__ == '__main__':
    sys.exit(main())
